from datetime import date
from datetime import timedelta
from io import BytesIO, StringIO
from types import SimpleNamespace
from unittest.mock import patch

from django.contrib.auth.models import Group, User
from django.core import mail
from django.core.management import call_command
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from .models import AgendamentoLab, Aluno, Laboratorio, Professor, RegistroFaltaAluno, RegistroOcorrenciaAluno, Turma
from .services.whatsapp_service import (
    enviar_mensagem_whatsapp,
    enviar_template_aviso_falta_aluno,
    enviar_template_aviso_ocorrencia_aluno,
    formatar_numero_whatsapp,
)
from .views.faltas import _telefone_whatsapp_aluno


class SegurancaSenhaTests(TestCase):
    def setUp(self):
        self.grupo_equipe = Group.objects.create(name='Equipe Diretiva')
        self.grupo_professores = Group.objects.create(name='Professores')

    def test_home_mostra_aviso_para_professor_com_senha_padrao(self):
        user = User.objects.create_user(username='professor-seguranca', password='zanchim2026')
        user.groups.add(self.grupo_professores)
        Professor.objects.create(usuario=user, nome_completo='Professor Seguranca', ativo=True)
        self.client.force_login(user)

        response = self.client.get(reverse('home'))

        self.assertContains(response, 'Por seguranca, altere sua senha pessoal.')
        self.assertContains(response, reverse('alterar_senha'))
        self.assertContains(response, reverse('recuperar_senha'))

    def test_painel_mostra_aviso_para_equipe_com_senha_padrao(self):
        user = User.objects.create_user(username='equipe-seguranca', password='zanchim2026')
        user.groups.add(self.grupo_equipe)
        self.client.force_login(user)

        response = self.client.get(reverse('painel_equipe'))

        self.assertContains(response, 'Por seguranca, altere sua senha pessoal.')
        self.assertContains(response, reverse('alterar_senha'))
        self.assertContains(response, reverse('recuperar_senha'))

    def test_aviso_nao_aparece_para_superuser(self):
        user = User.objects.create_superuser(username='admin-seguranca', password='zanchim2026')
        self.client.force_login(user)

        response = self.client.get(reverse('home'))

        self.assertNotContains(response, 'Por seguranca, altere sua senha pessoal.')

    def test_alterar_senha_funciona_para_usuario_logado(self):
        user = User.objects.create_user(username='troca-senha', password='zanchim2026')
        self.client.force_login(user)

        response = self.client.post(
            reverse('alterar_senha'),
            {
                'old_password': 'zanchim2026',
                'new_password1': 'SenhaForte2026!',
                'new_password2': 'SenhaForte2026!',
            },
        )

        self.assertRedirects(response, reverse('senha_alterada'))
        user.refresh_from_db()
        self.assertTrue(user.check_password('SenhaForte2026!'))

    @override_settings(EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend')
    def test_recuperacao_de_senha_envia_email_com_link(self):
        User.objects.create_user(
            username='recuperar-senha',
            email='professor@example.com',
            password='zanchim2026',
        )

        response = self.client.post(reverse('recuperar_senha'), {'email': 'professor@example.com'})

        self.assertRedirects(response, reverse('password_reset_done'))
        self.assertEqual(len(mail.outbox), 1)
        self.assertIn('/conta/redefinir/', mail.outbox[0].body)


class FaltasOcorrenciasSeparacaoTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username='digitador',
            password='senha-teste',
        )
        digitadores = Group.objects.create(name='Digitadores')
        self.user.groups.add(digitadores)
        equipe = Group.objects.create(name='Equipe Diretiva')
        self.equipe_user = User.objects.create_user(
            username='equipe',
            password='senha-teste',
        )
        self.equipe_user.groups.add(equipe)

        self.turma = Turma.objects.create(
            nome='1A',
            ano=2026,
            serie='1 Ano',
            ativa=True,
            turno='manha',
        )
        self.aluno = Aluno.objects.create(
            nome='Aluno Teste',
            numero=7,
            turma=self.turma,
        )

    def test_falta_aluno_fica_apenas_em_registro_falta_aluno(self):
        falta = RegistroFaltaAluno.objects.create(
            aluno=self.aluno,
            data=date(2026, 5, 15),
            quantidade_faltas=1,
            pedagoga='Sonia',
            registrado_por=self.user,
        )

        self.assertTrue(RegistroFaltaAluno.objects.filter(pk=falta.pk).exists())
        self.assertFalse(
            RegistroOcorrenciaAluno.objects.filter(
                aluno=self.aluno,
                data=falta.data,
            ).exists()
        )

    def test_ocorrencia_aluno_fica_apenas_em_registro_ocorrencia_aluno(self):
        ocorrencia = RegistroOcorrenciaAluno.objects.create(
            aluno=self.aluno,
            data=date(2026, 5, 15),
            tipo_ocorrencia='atraso',
            faltou=False,
            turno='manha',
            atendido_por='Sonia',
            registrado_por=self.user,
        )

        self.assertTrue(RegistroOcorrenciaAluno.objects.filter(pk=ocorrencia.pk).exists())
        self.assertFalse(
            RegistroFaltaAluno.objects.filter(
                aluno=self.aluno,
                data=ocorrencia.data,
            ).exists()
        )

    def test_fluxo_ocorrencia_digitador_rejeita_tipo_falta(self):
        self.client.force_login(self.user)

        response = self.client.post(
            reverse('formulario_digitador'),
            {
                'turma': self.turma.pk,
                'numero_aluno': self.aluno.numero,
                'nome_aluno': self.aluno.nome,
                'data': '2026-05-15',
                'turno': 'manha',
                'faltou': '',
                'tipo_ocorrencia': 'falta',
                'motivo_alegado': 'Teste de tipo invalido',
                'atendido_por': 'Sonia',
                'responsavel_contatado': '',
                'horario_chegada': '',
                'horario_contato': '',
                'alegado_responsavel': '',
            },
        )

        self.assertRedirects(response, reverse('formulario_digitador'))
        self.assertFalse(RegistroOcorrenciaAluno.objects.filter(aluno=self.aluno).exists())
        self.assertFalse(
            RegistroFaltaAluno.objects.filter(
                aluno=self.aluno,
                data=date(2026, 5, 15),
            ).exists()
        )

    def test_fluxo_ocorrencia_real_rejeita_tipo_falta_maiusculo(self):
        response = self._post_ocorrencia_registrar('Falta')

        self.assertRedirects(response, reverse('registrar_ocorrencia_aluno'))
        self.assertFalse(RegistroOcorrenciaAluno.objects.filter(aluno=self.aluno).exists())
        self.assertFalse(RegistroFaltaAluno.objects.filter(aluno=self.aluno).exists())

    def test_painel_nao_conta_faltas_historicas_como_ocorrencias(self):
        self.client.force_login(self.equipe_user)
        RegistroOcorrenciaAluno.objects.create(
            aluno=self.aluno,
            data=date(2026, 5, 15),
            tipo_ocorrencia='falta',
            faltou=True,
            turno='manha',
            atendido_por='Sonia',
            registrado_por=self.user,
        )
        RegistroOcorrenciaAluno.objects.create(
            aluno=self.aluno,
            data=date(2026, 5, 15),
            tipo_ocorrencia='Falta',
            faltou=True,
            turno='manha',
            atendido_por='Sonia',
            registrado_por=self.user,
        )
        RegistroOcorrenciaAluno.objects.create(
            aluno=self.aluno,
            data=date(2026, 5, 15),
            tipo_ocorrencia='atraso',
            faltou=False,
            turno='manha',
            atendido_por='Sonia',
            registrado_por=self.user,
        )

        response = self.client.get(reverse('painel_equipe'), {
            'dia': '15',
            'mes': '5',
            'ano': '2026',
            'dia_ocorrencias_ranking': '15',
            'mes_ocorrencias_ranking': '5',
            'ano_ocorrencias_ranking': '2026',
        })

        self.assertEqual(response.context['total_ocorrencias_periodo'], 1)
        self.assertEqual(response.context['ocorrencias_por_dia'], [1])
        self.assertEqual(response.context['ranking_tipos_ocorrencia'], [{'tipo': 'Atraso', 'total': 1}])
        self.assertEqual(sum(response.context['ocorrencias_totais_manha']), 1)
        self.assertNotIn('Falta', response.context['ranking_tipos_ocorrencia_labels_json'])
        self.assertNotIn('falta', response.context['ranking_tipos_ocorrencia_labels_json'])

    def test_conferencia_ocorrencias_nao_lista_falta_no_filtro_nem_na_tabela(self):
        self.client.force_login(self.equipe_user)
        RegistroOcorrenciaAluno.objects.create(
            aluno=self.aluno,
            data=date(2026, 5, 15),
            tipo_ocorrencia='falta',
            faltou=True,
            turno='manha',
            atendido_por='Sonia',
            registrado_por=self.user,
        )
        RegistroOcorrenciaAluno.objects.create(
            aluno=self.aluno,
            data=date(2026, 5, 15),
            tipo_ocorrencia='atraso',
            faltou=False,
            turno='manha',
            atendido_por='Sonia',
            registrado_por=self.user,
        )

        response = self.client.get(reverse('conferencia_ocorrencias'), {'data': '2026-05-15'})

        self.assertEqual(len(response.context['ocorrencias']), 1)
        self.assertEqual(response.context['ocorrencias'][0].tipo_ocorrencia, 'atraso')
        self.assertNotIn(('falta', 'Falta'), response.context['tipos_ocorrencia'])
        self.assertNotContains(response, 'value="falta"')

    def test_filtro_relatorio_ocorrencias_nao_oferece_falta(self):
        self.client.force_login(self.equipe_user)

        response = self.client.get(reverse('relatorio_filtro'))

        self.assertNotContains(response, 'value="falta"')

    def test_relatorio_ocorrencias_nao_exporta_faltas_historicas(self):
        self.client.force_login(self.equipe_user)
        RegistroOcorrenciaAluno.objects.create(
            aluno=self.aluno,
            data=date(2026, 5, 15),
            tipo_ocorrencia='Falta',
            faltou=True,
            turno='manha',
            atendido_por='Sonia',
            registrado_por=self.user,
        )
        RegistroOcorrenciaAluno.objects.create(
            aluno=self.aluno,
            data=date(2026, 5, 15),
            tipo_ocorrencia='atraso',
            faltou=False,
            turno='manha',
            atendido_por='Sonia',
            registrado_por=self.user,
        )

        response = self.client.get(reverse('relatorio_ocorrencias_alunos'), {'mes': '5', 'ano': '2026'})
        workbook = load_workbook(BytesIO(response.content))
        valores = [
            cell
            for row in workbook.active.iter_rows(values_only=True)
            for cell in row
            if cell is not None
        ]

        self.assertIn('Atraso', valores)
        self.assertNotIn('Falta', valores)

    def _post_registrar_falta(self, possui_atestado='nao', follow=True):
        self.client.force_login(self.user)
        return self.client.post(
            reverse('registrar_falta_aluno'),
            {
                'aluno': self.aluno.pk,
                'data': '2026-05-15',
                'quantidade': '1',
                'pedagoga': 'Sonia',
                'possui_atestado': possui_atestado,
                'observacoes': 'Teste de aviso',
            },
            follow=follow,
        )

    @patch('apps.views.faltas.enviar_template_aviso_falta_aluno')
    def test_falta_com_atestado_nao_envia_whatsapp(self, enviar_mock):
        response = self._post_registrar_falta(possui_atestado='sim')

        falta = RegistroFaltaAluno.objects.get(aluno=self.aluno)
        self.assertTrue(falta.justificada)
        enviar_mock.assert_not_called()
        self.assertContains(response, 'Falta registrada com atestado. WhatsApp nao enviado.')

    @patch('apps.views.faltas.enviar_template_aviso_falta_aluno')
    def test_falta_sem_telefone_salva_e_registra_alerta_sem_envio(self, enviar_mock):
        with self.assertLogs('apps.views.faltas', level='WARNING') as logs:
            response = self._post_registrar_falta(possui_atestado='nao')

        self.assertTrue(RegistroFaltaAluno.objects.filter(aluno=self.aluno).exists())
        enviar_mock.assert_not_called()
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Falta registrada, mas nao ha telefone cadastrado para envio.')
        self.assertIn('aluno sem telefone cadastrado', '\n'.join(logs.output))

    def test_telefone_whatsapp_usa_telefone_quando_responsavel_vazio(self):
        aluno = SimpleNamespace(telefone_responsavel='   ', telefone='5544999501967')

        self.assertEqual(_telefone_whatsapp_aluno(aluno), '5544999501967')

    @patch('apps.views.faltas.enviar_template_aviso_falta_aluno')
    def test_falta_com_falha_no_whatsapp_salva_e_registra_log(self, enviar_mock):
        self.aluno.telefone = '(44) 99999-0000'
        self.aluno.save(update_fields=['telefone'])
        enviar_mock.return_value = {
            'status': False,
            'numero': '5544999990000',
            'erro': {'message': 'Numero invalido'},
            'resposta': {'error': {'message': 'Numero invalido'}},
            'status_code': 400,
        }

        with self.assertLogs('apps.views.faltas', level='WARNING') as logs:
            response = self._post_registrar_falta(possui_atestado='nao')

        self.assertTrue(RegistroFaltaAluno.objects.filter(aluno=self.aluno).exists())
        enviar_mock.assert_called_once()
        self.assertEqual(enviar_mock.call_args.args[0], '(44) 99999-0000')
        self.assertEqual(enviar_mock.call_args.args[1], self.aluno.nome)
        self.assertEqual(enviar_mock.call_args.args[2], '15/05/2026')
        self.assertContains(response, 'Falta registrada, mas o aviso nao foi enviado. Confira o numero do responsavel.')
        self.assertIn('envio de WhatsApp falhou', '\n'.join(logs.output))

    @patch('apps.views.faltas.enviar_template_aviso_falta_aluno')
    def test_falta_com_whatsapp_enviado_exibe_sucesso(self, enviar_mock):
        self.aluno.telefone = '(44) 99999-0000'
        self.aluno.save(update_fields=['telefone'])
        enviar_mock.return_value = {
            'status': True,
            'numero': '5544999990000',
            'resposta': {'messages': [{'id': 'wamid.teste'}]},
            'erro': None,
            'status_code': 200,
        }

        response = self._post_registrar_falta(possui_atestado='nao')

        self.assertTrue(RegistroFaltaAluno.objects.filter(aluno=self.aluno).exists())
        enviar_mock.assert_called_once()
        self.assertEqual(enviar_mock.call_args.args[0], '(44) 99999-0000')
        self.assertEqual(enviar_mock.call_args.args[1], self.aluno.nome)
        self.assertEqual(enviar_mock.call_args.args[2], '15/05/2026')
        self.assertContains(response, 'Falta registrada e aviso enviado pelo WhatsApp.')

    @patch('apps.views.alunos.enviar_template_aviso_ocorrencia_aluno')
    def test_ocorrencia_com_whatsapp_enviado_exibe_sucesso(self, enviar_mock):
        self.client.force_login(self.user)
        self.aluno.telefone = '(44) 99999-0000'
        self.aluno.save(update_fields=['telefone'])
        enviar_mock.return_value = {
            'status': True,
            'numero': '5544999990000',
            'resposta': {'messages': [{'id': 'wamid.teste'}]},
            'erro': None,
            'status_code': 200,
        }

        response = self.client.post(
            reverse('formulario_digitador'),
            {
                'turma': self.turma.pk,
                'numero_aluno': self.aluno.numero,
                'nome_aluno': self.aluno.nome,
                'data': '2026-05-15',
                'turno': 'manha',
                'faltou': '',
                'tipo_ocorrencia': 'piercing',
                'motivo_alegado': 'Uso de piercing',
                'atendido_por': 'Sonia',
                'responsavel_contatado': '',
                'horario_chegada': '',
                'horario_contato': '',
                'alegado_responsavel': '',
            },
            follow=True,
        )

        self.assertTrue(RegistroOcorrenciaAluno.objects.filter(aluno=self.aluno, tipo_ocorrencia='piercing').exists())
        enviar_mock.assert_called_once()
        self.assertEqual(enviar_mock.call_args.args[0], '(44) 99999-0000')
        self.assertEqual(enviar_mock.call_args.args[1], self.aluno.nome)
        self.assertEqual(enviar_mock.call_args.args[2], 'Uso de Piercing')
        self.assertEqual(enviar_mock.call_args.args[3], '15/05/2026')
        self.assertContains(response, 'Ocorrencia registrada e aviso enviado pelo WhatsApp.')

    @patch('apps.views.alunos.enviar_template_aviso_ocorrencia_aluno')
    def test_ocorrencia_sem_telefone_salva_e_alerta_sem_envio(self, enviar_mock):
        self.client.force_login(self.user)

        with self.assertLogs('apps.views.alunos', level='WARNING') as logs:
            response = self.client.post(
                reverse('formulario_digitador'),
                {
                    'turma': self.turma.pk,
                    'numero_aluno': self.aluno.numero,
                    'nome_aluno': self.aluno.nome,
                    'data': '2026-05-15',
                    'turno': 'manha',
                    'faltou': '',
                    'tipo_ocorrencia': 'uniforme',
                    'motivo_alegado': 'Uniforme',
                    'atendido_por': 'Sonia',
                    'responsavel_contatado': '',
                    'horario_chegada': '',
                    'horario_contato': '',
                    'alegado_responsavel': '',
                },
                follow=True,
            )

        self.assertTrue(RegistroOcorrenciaAluno.objects.filter(aluno=self.aluno, tipo_ocorrencia='uniforme').exists())
        enviar_mock.assert_not_called()
        self.assertContains(response, 'Ocorrencia registrada, mas nao ha telefone cadastrado.')
        self.assertIn('aluno sem telefone cadastrado', '\n'.join(logs.output))

    def _post_ocorrencia_digitador(self, tipo_ocorrencia, follow=True, enviar_reincidencia=''):
        self.client.force_login(self.user)
        return self.client.post(
            reverse('formulario_digitador'),
            {
                'turma': self.turma.pk,
                'numero_aluno': self.aluno.numero,
                'nome_aluno': self.aluno.nome,
                'data': '2026-05-15',
                'turno': 'manha',
                'faltou': '',
                'tipo_ocorrencia': tipo_ocorrencia,
                'motivo_alegado': 'Teste de ocorrencia',
                'atendido_por': 'Sonia',
                'responsavel_contatado': '',
                'horario_chegada': '',
                'horario_contato': '',
                'alegado_responsavel': '',
                'enviar_whatsapp_reincidencia': enviar_reincidencia,
            },
            follow=follow,
        )

    def _post_ocorrencia_registrar(self, tipo_ocorrencia, follow=True, enviar_reincidencia=''):
        self.client.force_login(self.user)
        return self.client.post(
            reverse('registrar_ocorrencia_aluno'),
            {
                'pedagoga': 'Sonia',
                'turma': self.turma.pk,
                'numero_aluno': self.aluno.numero,
                'nome_aluno': self.aluno.nome,
                'data': '2026-05-15',
                'turno': 'manha',
                'faltou': '',
                'tipo_ocorrencia': tipo_ocorrencia,
                'motivo_alegado': 'Teste de ocorrencia',
                'responsavel_contatado': '',
                'horario_chegada': '',
                'horario_contato': '',
                'alegado_responsavel': '',
                'observacoes_adicionais': '',
                'enviar_whatsapp_reincidencia': enviar_reincidencia,
            },
            follow=follow,
        )

    @patch('apps.views.faltas.enviar_template_aviso_ocorrencia_aluno')
    def test_ocorrencia_atraso_rota_real_envia_whatsapp(self, enviar_mock):
        self.aluno.telefone = '(44) 99999-0000'
        self.aluno.save(update_fields=['telefone'])
        enviar_mock.return_value = {
            'status': True,
            'numero': '5544999990000',
            'resposta': {'messages': [{'id': 'wamid.teste'}]},
            'erro': None,
            'status_code': 200,
        }

        response = self._post_ocorrencia_registrar(' atraso ')

        self.assertTrue(RegistroOcorrenciaAluno.objects.filter(aluno=self.aluno, tipo_ocorrencia='atraso').exists())
        enviar_mock.assert_called_once()
        self.assertEqual(enviar_mock.call_args.args[0], '(44) 99999-0000')
        self.assertEqual(enviar_mock.call_args.args[1], self.aluno.nome)
        self.assertEqual(enviar_mock.call_args.args[2], 'Atraso')
        self.assertEqual(enviar_mock.call_args.args[3], '15/05/2026')
        self.assertContains(response, 'Ocorrencia registrada e aviso enviado pelo WhatsApp.')

    @patch('apps.views.faltas.enviar_template_aviso_ocorrencia_aluno')
    def test_ocorrencia_atraso_rota_real_segunda_vez_nao_reenvia(self, enviar_mock):
        self.aluno.telefone = '(44) 99999-0000'
        self.aluno.save(update_fields=['telefone'])
        RegistroOcorrenciaAluno.objects.create(
            aluno=self.aluno,
            data=date(2026, 5, 15),
            tipo_ocorrencia='atraso',
            faltou=False,
            turno='manha',
            atendido_por='Sonia',
            registrado_por=self.user,
        )

        response = self._post_ocorrencia_registrar('atraso', enviar_reincidencia='sim')

        enviar_mock.assert_not_called()
        self.assertContains(response, 'Ja existe um registro para este aluno nesta data com o mesmo tipo de ocorrencia.')
        self.assertEqual(RegistroOcorrenciaAluno.objects.filter(aluno=self.aluno, tipo_ocorrencia='atraso').count(), 1)

    @patch('apps.views.faltas.enviar_template_aviso_ocorrencia_aluno')
    def test_ocorrencia_piercing_rota_real_reincidencia_reenvia(self, enviar_mock):
        self.aluno.telefone = '(44) 99999-0000'
        self.aluno.save(update_fields=['telefone'])
        RegistroOcorrenciaAluno.objects.create(
            aluno=self.aluno,
            data=date(2026, 5, 15),
            tipo_ocorrencia='piercing',
            faltou=False,
            turno='manha',
            atendido_por='Sonia',
            registrado_por=self.user,
        )
        enviar_mock.return_value = {
            'status': True,
            'numero': '5544999990000',
            'resposta': {'messages': [{'id': 'wamid.teste'}]},
            'erro': None,
            'status_code': 200,
        }

        response = self._post_ocorrencia_registrar('piercing', enviar_reincidencia='sim')

        enviar_mock.assert_called_once()
        self.assertEqual(enviar_mock.call_args.args[2], 'Uso de Piercing')
        self.assertContains(response, 'Novo aviso de reincidencia enviado pelo WhatsApp.')
        self.assertEqual(RegistroOcorrenciaAluno.objects.filter(aluno=self.aluno, tipo_ocorrencia='piercing').count(), 1)

    @patch('apps.views.alunos.enviar_template_aviso_ocorrencia_aluno')
    def test_ocorrencia_atraso_primeira_vez_envia_whatsapp(self, enviar_mock):
        self.aluno.telefone = '(44) 99999-0000'
        self.aluno.save(update_fields=['telefone'])
        enviar_mock.return_value = {
            'status': True,
            'numero': '5544999990000',
            'resposta': {'messages': [{'id': 'wamid.teste'}]},
            'erro': None,
            'status_code': 200,
        }

        response = self._post_ocorrencia_digitador('atraso')

        self.assertTrue(RegistroOcorrenciaAluno.objects.filter(aluno=self.aluno, tipo_ocorrencia='atraso').exists())
        enviar_mock.assert_called_once()
        self.assertEqual(enviar_mock.call_args.args[0], '(44) 99999-0000')
        self.assertEqual(enviar_mock.call_args.args[1], self.aluno.nome)
        self.assertEqual(enviar_mock.call_args.args[2], 'Atraso')
        self.assertEqual(enviar_mock.call_args.args[3], '15/05/2026')
        self.assertContains(response, 'Ocorrencia registrada e aviso enviado pelo WhatsApp.')

    @patch('apps.views.alunos.enviar_template_aviso_ocorrencia_aluno')
    def test_ocorrencia_atraso_com_maiuscula_envia_whatsapp(self, enviar_mock):
        self.aluno.telefone = '(44) 99999-0000'
        self.aluno.save(update_fields=['telefone'])
        enviar_mock.return_value = {
            'status': True,
            'numero': '5544999990000',
            'resposta': {'messages': [{'id': 'wamid.teste'}]},
            'erro': None,
            'status_code': 200,
        }

        response = self._post_ocorrencia_digitador('Atraso')

        self.assertTrue(RegistroOcorrenciaAluno.objects.filter(aluno=self.aluno, tipo_ocorrencia='atraso').exists())
        enviar_mock.assert_called_once()
        self.assertEqual(enviar_mock.call_args.args[2], 'Atraso')
        self.assertContains(response, 'Ocorrencia registrada e aviso enviado pelo WhatsApp.')

    @patch('apps.views.alunos.enviar_template_aviso_ocorrencia_aluno')
    def test_ocorrencia_atraso_com_espacos_envia_whatsapp(self, enviar_mock):
        self.aluno.telefone = '(44) 99999-0000'
        self.aluno.save(update_fields=['telefone'])
        enviar_mock.return_value = {
            'status': True,
            'numero': '5544999990000',
            'resposta': {'messages': [{'id': 'wamid.teste'}]},
            'erro': None,
            'status_code': 200,
        }

        response = self._post_ocorrencia_digitador(' atraso ')

        self.assertTrue(RegistroOcorrenciaAluno.objects.filter(aluno=self.aluno, tipo_ocorrencia='atraso').exists())
        enviar_mock.assert_called_once()
        self.assertEqual(enviar_mock.call_args.args[2], 'Atraso')
        self.assertContains(response, 'Ocorrencia registrada e aviso enviado pelo WhatsApp.')

    @patch('apps.views.alunos.enviar_template_aviso_ocorrencia_aluno')
    def test_ocorrencia_atraso_segunda_vez_nao_entra_reincidencia(self, enviar_mock):
        self.aluno.telefone = '(44) 99999-0000'
        self.aluno.save(update_fields=['telefone'])
        RegistroOcorrenciaAluno.objects.create(
            aluno=self.aluno,
            data=date(2026, 5, 15),
            tipo_ocorrencia='atraso',
            faltou=False,
            turno='manha',
            atendido_por='Sonia',
            registrado_por=self.user,
        )

        response = self._post_ocorrencia_digitador('atraso', enviar_reincidencia='sim')

        enviar_mock.assert_not_called()
        self.assertContains(response, 'Ja existe um registro para este aluno nesta data com o mesmo tipo de ocorrencia.')
        self.assertEqual(RegistroOcorrenciaAluno.objects.filter(aluno=self.aluno, tipo_ocorrencia='atraso').count(), 1)

    @patch('apps.views.alunos.enviar_template_aviso_ocorrencia_aluno')
    def test_ocorrencia_piercing_continua_com_reincidencia(self, enviar_mock):
        self.aluno.telefone = '(44) 99999-0000'
        self.aluno.save(update_fields=['telefone'])
        RegistroOcorrenciaAluno.objects.create(
            aluno=self.aluno,
            data=date(2026, 5, 15),
            tipo_ocorrencia='piercing',
            faltou=False,
            turno='manha',
            atendido_por='Sonia',
            registrado_por=self.user,
        )
        enviar_mock.return_value = {
            'status': True,
            'numero': '5544999990000',
            'resposta': {'messages': [{'id': 'wamid.teste'}]},
            'erro': None,
            'status_code': 200,
        }

        response = self._post_ocorrencia_digitador('piercing', enviar_reincidencia='sim')

        enviar_mock.assert_called_once()
        self.assertEqual(enviar_mock.call_args.args[2], 'Uso de Piercing')
        self.assertContains(response, 'Novo aviso de reincidencia enviado pelo WhatsApp.')
        self.assertEqual(RegistroOcorrenciaAluno.objects.filter(aluno=self.aluno, tipo_ocorrencia='piercing').count(), 1)

    def test_comando_migrar_faltas_ocorrencias_dry_run_nao_grava(self):
        RegistroOcorrenciaAluno.objects.create(
            aluno=self.aluno,
            data=date(2026, 5, 16),
            tipo_ocorrencia='falta',
            faltou=True,
            turno='manha',
            atendido_por='Sonia',
            motivo_alegado='Ausencia informada',
            observacoes_adicionais='Observacao antiga',
            registrado_por=self.user,
        )
        out = StringIO()

        call_command('migrar_faltas_ocorrencias', stdout=out)

        self.assertEqual(RegistroFaltaAluno.objects.count(), 0)
        saida = out.getvalue()
        self.assertIn('Modo: DRY-RUN', saida)
        self.assertIn('Total encontrado: 1', saida)
        self.assertIn('Total que sera migrado: 1', saida)
        self.assertIn('Duplicados ignorados: 0', saida)

    def test_comando_migrar_faltas_ocorrencias_confirma_e_ignora_duplicados(self):
        outro_aluno = Aluno.objects.create(nome='Outro Aluno', numero=8, turma=self.turma)
        RegistroFaltaAluno.objects.create(
            aluno=self.aluno,
            data=date(2026, 5, 16),
            quantidade_faltas=1,
            pedagoga='Sonia',
            registrado_por=self.user,
        )
        RegistroOcorrenciaAluno.objects.create(
            aluno=self.aluno,
            data=date(2026, 5, 16),
            tipo_ocorrencia='falta',
            faltou=True,
            turno='manha',
            atendido_por='Sonia',
            registrado_por=self.user,
        )
        RegistroOcorrenciaAluno.objects.create(
            aluno=outro_aluno,
            data=date(2026, 5, 17),
            tipo_ocorrencia='Falta',
            faltou=True,
            turno='manha',
            atendido_por='Simone',
            motivo_alegado='Faltou no periodo',
            observacoes_adicionais='Contato pendente',
            registrado_por=self.user,
        )
        out = StringIO()

        call_command('migrar_faltas_ocorrencias', '--confirmar', stdout=out)

        self.assertEqual(RegistroFaltaAluno.objects.count(), 2)
        falta_migrada = RegistroFaltaAluno.objects.get(aluno=outro_aluno)
        self.assertEqual(falta_migrada.data, date(2026, 5, 17))
        self.assertEqual(falta_migrada.pedagoga, 'Simone')
        self.assertEqual(falta_migrada.registrado_por, self.user)
        self.assertIn('Contato pendente', falta_migrada.observacoes)
        self.assertIn('Faltou no periodo', falta_migrada.observacoes)
        self.assertEqual(RegistroOcorrenciaAluno.objects.filter(tipo_ocorrencia__in=['falta', 'Falta']).count(), 2)

        saida = out.getvalue()
        self.assertIn('Total encontrado: 2', saida)
        self.assertIn('Total que sera migrado: 1', saida)
        self.assertIn('Duplicados ignorados: 1', saida)
        self.assertIn('Total migrado: 1', saida)


class WhatsAppServiceTests(TestCase):
    def test_formatar_numero_whatsapp_adiciona_codigo_do_brasil(self):
        self.assertEqual(formatar_numero_whatsapp('(44) 99999-0000'), '5544999990000')

    @override_settings(WHATSAPP_TOKEN='', PHONE_NUMBER_ID='')
    def test_enviar_mensagem_whatsapp_sem_configuracao_retorna_erro_controlado(self):
        resultado = enviar_mensagem_whatsapp('(44) 99999-0000', 'Mensagem de teste')

        self.assertFalse(resultado['status'])
        self.assertEqual(resultado['numero'], '5544999990000')
        self.assertIn('WHATSAPP_TOKEN', resultado['erro'])

    @override_settings(
        WHATSAPP_TOKEN='token-teste',
        PHONE_NUMBER_ID='123456',
        WHATSAPP_API_VERSION='v20.0',
    )
    @patch('apps.services.whatsapp_service.requests.post')
    def test_enviar_template_aviso_falta_aluno_usa_template_aprovado(self, post_mock):
        post_mock.return_value.ok = True
        post_mock.return_value.status_code = 200
        post_mock.return_value.json.return_value = {'messages': [{'id': 'wamid.teste'}]}

        resultado = enviar_template_aviso_falta_aluno('44999990000', 'Aluno Teste', '15/05/2026')

        self.assertTrue(resultado['status'])
        payload = post_mock.call_args.kwargs['json']
        self.assertEqual(payload['type'], 'template')
        self.assertEqual(payload['template']['name'], 'aviso_falta_aluno')
        self.assertEqual(payload['template']['language']['code'], 'pt_BR')
        parametros = payload['template']['components'][0]['parameters']
        self.assertEqual(parametros[0]['text'], 'Aluno Teste')
        self.assertEqual(parametros[1]['text'], '15/05/2026')

    @override_settings(
        WHATSAPP_TOKEN='token-teste',
        PHONE_NUMBER_ID='123456',
        WHATSAPP_API_VERSION='v20.0',
    )
    @patch('apps.services.whatsapp_service.requests.post')
    def test_enviar_template_aviso_ocorrencia_aluno_usa_template_aprovado(self, post_mock):
        post_mock.return_value.ok = True
        post_mock.return_value.status_code = 200
        post_mock.return_value.json.return_value = {'messages': [{'id': 'wamid.teste'}]}

        resultado = enviar_template_aviso_ocorrencia_aluno(
            '44999990000',
            'Aluno Teste',
            'Uso de Piercing',
            '15/05/2026',
        )

        self.assertTrue(resultado['status'])
        payload = post_mock.call_args.kwargs['json']
        self.assertEqual(payload['type'], 'template')
        self.assertEqual(payload['template']['name'], 'aviso_ocorrencia_aluno')
        self.assertEqual(payload['template']['language']['code'], 'pt_BR')
        parametros = payload['template']['components'][0]['parameters']
        self.assertEqual(parametros[0]['text'], 'Aluno Teste')
        self.assertEqual(parametros[1]['text'], 'Uso de Piercing')
        self.assertEqual(parametros[2]['text'], '15/05/2026')


class CronogramaLaboratoriosTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='diretivo', password='senha-teste')
        self.client.force_login(self.user)
        self.turma = Turma.objects.create(nome='2A', ano=2026, serie='2 Ano', ativa=True, turno='manha')
        self.professor_manha = Professor.objects.create(nome_completo='Ana Maria Silva', ativo=True)
        self.professor_tarde = Professor.objects.create(nome_completo='Bruno Costa Lima', ativo=True)
        self.labs = [
            Laboratorio.objects.create(nome=f'Laboratorio {i}', tipo='fixo', equipamento='Computadores', ativo=True)
            for i in range(1, 6)
        ]
        self.data_segunda = date(2026, 5, 18)
        AgendamentoLab.objects.create(
            laboratorio=self.labs[0],
            data=self.data_segunda,
            horario='1',
            turno='manha',
            professor=self.professor_manha,
            turma=self.turma,
            disciplina='Matematica',
            observacao='Levar projetor',
            registrado_por=self.user,
        )
        AgendamentoLab.objects.create(
            laboratorio=self.labs[1],
            data=self.data_segunda,
            horario='2',
            turno='tarde',
            professor=self.professor_tarde,
            turma=self.turma,
            disciplina='Historia',
            registrado_por=self.user,
        )

    def _get_cronograma(self, **params):
        base = {'data_inicio': self.data_segunda.isoformat()}
        base.update(params)
        return self.client.get(reverse('cronograma_semanal'), base)

    def test_cronograma_filtra_por_laboratorio(self):
        response = self._get_cronograma(laboratorio=str(self.labs[0].id))

        self.assertEqual(list(response.context['laboratorios']), [self.labs[0]])
        self.assertContains(response, 'Matematica')
        self.assertNotContains(response, 'Historia')

    def test_cronograma_filtra_turno_manha(self):
        response = self._get_cronograma(turno='manha')

        self.assertContains(response, 'Matematica')
        self.assertNotContains(response, 'Historia')
        self.assertNotContains(response, 'data-turno-separator="manha"')
        self.assertNotContains(response, 'data-turno-separator="tarde"')
        self.assertTrue(all(slot['turno'] == 'manha' for slot in response.context['slots_cronograma']))

    def test_cronograma_filtra_turno_tarde(self):
        response = self._get_cronograma(turno='tarde')

        self.assertContains(response, 'Historia')
        self.assertNotContains(response, 'Matematica')
        self.assertNotContains(response, 'data-turno-separator="manha"')
        self.assertNotContains(response, 'data-turno-separator="tarde"')
        self.assertTrue(all(slot['turno'] == 'tarde' for slot in response.context['slots_cronograma']))

    def test_cronograma_filtra_laboratorio_e_turno(self):
        response = self._get_cronograma(laboratorio=str(self.labs[1].id), turno='tarde')

        self.assertEqual(list(response.context['laboratorios']), [self.labs[1]])
        self.assertContains(response, 'Historia')
        self.assertNotContains(response, 'Matematica')

    def test_cronograma_todos_laboratorios(self):
        response = self._get_cronograma()

        self.assertEqual(len(response.context['laboratorios']), 5)
        for lab in self.labs:
            self.assertContains(response, lab.nome)

    def test_cronograma_todos_turnos(self):
        response = self._get_cronograma()

        turnos = {slot['turno'] for slot in response.context['slots_cronograma']}
        self.assertEqual(turnos, {'manha', 'tarde'})
        self.assertContains(response, 'Matematica')
        self.assertContains(response, 'Historia')
        self.assertContains(response, 'data-turno-separator="manha"')
        self.assertContains(response, 'data-turno-separator="tarde"')

    def test_cronograma_nao_mostra_botoes_de_navegacao_semanal(self):
        response = self._get_cronograma()

        self.assertNotContains(response, 'Semana Anterior')
        self.assertNotContains(response, 'Proxima Semana')

    def test_cronograma_sem_data_inicio_usa_semana_mais_recente_com_agendamentos(self):
        data_recente = date(2026, 6, 3)
        AgendamentoLab.objects.create(
            laboratorio=self.labs[2],
            data=data_recente,
            horario='3',
            turno='manha',
            professor=self.professor_manha,
            turma=self.turma,
            disciplina='Robotica',
            registrado_por=self.user,
        )

        response = self.client.get(reverse('cronograma_semanal'))

        self.assertEqual(response.context['data_inicio'], date(2026, 6, 1))
        self.assertContains(response, 'Robotica')
        self.assertNotContains(response, 'Matematica')

    def test_cronograma_professores_aparecem_no_dia_correto(self):
        response = self._get_cronograma()

        self.assertContains(response, 'Segunda-feira')
        self.assertContains(response, self.professor_manha.nome_abreviado)
        self.assertContains(response, self.professor_tarde.nome_abreviado)
        self.assertContains(response, 'Levar projetor')


class HomeProfessorLaboratoriosTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='professor', password='senha-teste')
        self.outro_user = User.objects.create_user(username='outro-professor', password='senha-teste')
        self.professor = Professor.objects.create(
            usuario=self.user,
            nome_completo='Carlos Alberto Souza',
            ativo=True,
        )
        self.outro_professor = Professor.objects.create(
            usuario=self.outro_user,
            nome_completo='Marina Oliveira',
            ativo=True,
        )
        self.turma = Turma.objects.create(nome='3A', ano=2026, serie='3 Ano', ativa=True, turno='manha')
        self.laboratorio_1 = Laboratorio.objects.create(nome='Laboratorio 1', tipo='fixo', equipamento='Computadores', ativo=True)
        self.laboratorio_2 = Laboratorio.objects.create(nome='Laboratorio 2', tipo='fixo', equipamento='Computadores', ativo=True)
        self.semana_atual_inicio = timezone.localdate() - timedelta(days=timezone.localdate().weekday())

    def _agendar(self, professor, laboratorio, data, horario='1', disciplina='Fisica', observacao='', turno='manha'):
        return AgendamentoLab.objects.create(
            laboratorio=laboratorio,
            data=data,
            horario=horario,
            turno=turno,
            professor=professor,
            turma=self.turma,
            disciplina=disciplina,
            observacao=observacao,
            registrado_por=self.user,
        )

    def test_home_mostra_link_discreto_para_meus_horarios(self):
        self.client.force_login(self.user)

        response = self.client.get(reverse('home'))

        self.assertContains(response, 'Meus Horarios de Laboratorio')
        self.assertContains(response, reverse('meus_horarios_laboratorio'))
        self.assertNotContains(response, 'Nenhum horario de laboratorio encontrado.')

    def test_professor_logado_ve_apenas_proprios_agendamentos(self):
        self._agendar(
            self.professor,
            self.laboratorio_1,
            self.semana_atual_inicio,
            horario='1',
            disciplina='Fisica',
            observacao='Levar roteiro',
        )
        self._agendar(
            self.outro_professor,
            self.laboratorio_2,
            self.semana_atual_inicio,
            horario='2',
            disciplina='Quimica',
        )
        self.client.force_login(self.user)

        response = self.client.get(reverse('meus_horarios_laboratorio'))

        self.assertContains(response, 'Meus Horarios de Laboratorio')
        self.assertContains(response, 'Laboratorio 1')
        self.assertContains(response, 'Fisica')
        self.assertContains(response, 'Levar roteiro')
        self.assertNotContains(response, 'Laboratorio 2')
        self.assertNotContains(response, 'Quimica')
        self.assertEqual(len(response.context['horarios_laboratorio_professor']['agendamentos']), 1)

    def test_meus_horarios_filtro_manha_funciona(self):
        self._agendar(self.professor, self.laboratorio_1, self.semana_atual_inicio, disciplina='Fisica', turno='manha')
        self._agendar(self.professor, self.laboratorio_2, self.semana_atual_inicio, horario='2', disciplina='Robotica', turno='tarde')
        self.client.force_login(self.user)

        response = self.client.get(reverse('meus_horarios_laboratorio'), {'turno': 'manha'})

        self.assertContains(response, 'Fisica')
        self.assertContains(response, 'Manha')
        self.assertNotContains(response, 'Robotica')
        self.assertNotContains(response, 'data-turno-separator="manha"')
        self.assertNotContains(response, 'data-turno-separator="tarde"')
        self.assertEqual(response.context['horarios_laboratorio_professor']['turno_filtro'], 'manha')
        self.assertEqual(len(response.context['horarios_laboratorio_professor']['agendamentos']), 1)

    def test_meus_horarios_filtro_tarde_funciona(self):
        self._agendar(self.professor, self.laboratorio_1, self.semana_atual_inicio, disciplina='Fisica', turno='manha')
        self._agendar(self.professor, self.laboratorio_2, self.semana_atual_inicio, horario='2', disciplina='Robotica', turno='tarde')
        self.client.force_login(self.user)

        response = self.client.get(reverse('meus_horarios_laboratorio'), {'turno': 'tarde'})

        self.assertContains(response, 'Robotica')
        self.assertContains(response, 'Tarde')
        self.assertNotContains(response, 'Fisica')
        self.assertNotContains(response, 'data-turno-separator="manha"')
        self.assertNotContains(response, 'data-turno-separator="tarde"')
        self.assertEqual(response.context['horarios_laboratorio_professor']['turno_filtro'], 'tarde')
        self.assertEqual(len(response.context['horarios_laboratorio_professor']['agendamentos']), 1)

    def test_meus_horarios_filtro_todos_funciona(self):
        self._agendar(self.professor, self.laboratorio_1, self.semana_atual_inicio, disciplina='Fisica', turno='manha')
        self._agendar(self.professor, self.laboratorio_2, self.semana_atual_inicio, horario='2', disciplina='Robotica', turno='tarde')
        self.client.force_login(self.user)

        response = self.client.get(reverse('meus_horarios_laboratorio'))

        self.assertContains(response, 'Fisica')
        self.assertContains(response, 'Robotica')
        self.assertContains(response, 'data-turno-separator="manha"')
        self.assertContains(response, 'data-turno-separator="tarde"')
        conteudo = response.content.decode()
        self.assertLess(conteudo.index('MANHA'), conteudo.index('Fisica'))
        self.assertLess(conteudo.index('Fisica'), conteudo.index('TARDE'))
        self.assertLess(conteudo.index('TARDE'), conteudo.index('Robotica'))
        self.assertEqual(response.context['horarios_laboratorio_professor']['turno_filtro'], '')
        self.assertEqual(len(response.context['horarios_laboratorio_professor']['agendamentos']), 2)

    def test_meus_horarios_filtro_sem_resultado_mostra_mensagem_do_turno(self):
        self._agendar(self.professor, self.laboratorio_1, self.semana_atual_inicio, disciplina='Fisica', turno='manha')
        self.client.force_login(self.user)

        response = self.client.get(reverse('meus_horarios_laboratorio'), {'turno': 'tarde'})

        self.assertContains(response, 'Nenhum horario encontrado para este turno.')
        self.assertNotContains(response, 'Fisica')
        self.assertEqual(response.context['horarios_laboratorio_professor']['agendamentos'], [])

    def test_meus_horarios_usa_semana_atual_quando_tem_agendamento(self):
        data_atual = self.semana_atual_inicio + timedelta(days=1)
        data_antiga = self.semana_atual_inicio - timedelta(days=14)
        self._agendar(self.professor, self.laboratorio_1, data_atual, disciplina='Biologia')
        self._agendar(self.professor, self.laboratorio_2, data_antiga, horario='2', disciplina='Geografia')
        self.client.force_login(self.user)

        response = self.client.get(reverse('meus_horarios_laboratorio'))

        self.assertContains(response, 'Biologia')
        self.assertNotContains(response, 'Geografia')
        self.assertEqual(response.context['horarios_laboratorio_professor']['semana_inicio'], self.semana_atual_inicio)

    def test_meus_horarios_sem_semana_atual_pega_semana_mais_recente(self):
        data_antiga = self.semana_atual_inicio - timedelta(days=14)
        data_recente = self.semana_atual_inicio - timedelta(days=7)
        self._agendar(self.professor, self.laboratorio_1, data_antiga, disciplina='Artes')
        self._agendar(self.professor, self.laboratorio_2, data_recente, horario='2', disciplina='Robotica')
        self.client.force_login(self.user)

        response = self.client.get(reverse('meus_horarios_laboratorio'))

        self.assertContains(response, 'Robotica')
        self.assertNotContains(response, 'Artes')
        self.assertEqual(
            response.context['horarios_laboratorio_professor']['semana_inicio'],
            data_recente - timedelta(days=data_recente.weekday()),
        )

    def test_meus_horarios_sem_agendamento_mostra_mensagem(self):
        self.client.force_login(self.user)

        response = self.client.get(reverse('meus_horarios_laboratorio'))

        self.assertContains(response, 'Nenhum horario de laboratorio encontrado.')
        self.assertEqual(response.context['horarios_laboratorio_professor']['agendamentos'], [])
