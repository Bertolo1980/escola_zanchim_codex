from .utilitarios import *

import logging

from django.db import transaction

from apps.forms import TIPOS_OCORRENCIA_PERMITIDOS, normalizar_tipo_ocorrencia
from apps.services.whatsapp_service import enviar_template_aviso_ocorrencia_aluno

logger = logging.getLogger(__name__)

TIPOS_OCORRENCIA_WHATSAPP = {'atraso', 'piercing', 'desvio_normas', 'cabelo', 'uniforme', 'fora_sala', 'matando_aula'}
TIPOS_OCORRENCIA_REINCIDENCIA = {'piercing', 'desvio_normas', 'cabelo', 'uniforme'}
TIPOS_OCORRENCIA_LABELS = {
    'atraso': 'Atraso',
    'piercing': 'Uso de Piercing',
    'desvio_normas': 'Desvio de Normas',
    'cabelo': 'Cabelo',
    'uniforme': 'Uniforme',
    'fora_sala': 'Fora da sala',
    'matando_aula': 'Matando aula',
}


def _telefone_whatsapp_ocorrencia(aluno):
    for campo in ('telefone_responsavel', 'telefone'):
        telefone = getattr(aluno, campo, None)
        if telefone is None:
            continue
        telefone = str(telefone).strip()
        if telefone:
            return telefone
    return ''


def _data_ocorrencia_whatsapp(data_ocorrencia):
    try:
        return datetime.strptime(str(data_ocorrencia), '%Y-%m-%d').strftime('%d/%m/%Y')
    except ValueError:
        return str(data_ocorrencia)


def _enviar_aviso_ocorrencia_whatsapp_digitador(request, aluno, tipo_ocorrencia, data_ocorrencia, mensagem_sucesso=None):
    tipo_ocorrencia = (tipo_ocorrencia or '').strip().lower()
    if tipo_ocorrencia == 'falta' or tipo_ocorrencia not in TIPOS_OCORRENCIA_WHATSAPP:
        return

    telefone = _telefone_whatsapp_ocorrencia(aluno)
    if not telefone:
        logger.warning(
            'Ocorrencia registrada sem envio de WhatsApp: aluno sem telefone cadastrado.',
            extra={
                'aluno': aluno.nome,
                'aluno_id': aluno.id,
                'tipo_ocorrencia': tipo_ocorrencia,
                'data_ocorrencia': str(data_ocorrencia),
            },
        )
        messages.warning(request, 'Ocorrencia registrada, mas nao ha telefone cadastrado. Confira o telefone do aluno/responsavel.')
        return

    resultado_whatsapp = enviar_template_aviso_ocorrencia_aluno(
        telefone,
        aluno.nome,
        TIPOS_OCORRENCIA_LABELS.get(tipo_ocorrencia, tipo_ocorrencia),
        _data_ocorrencia_whatsapp(data_ocorrencia),
    )

    if resultado_whatsapp.get('status'):
        messages.success(request, mensagem_sucesso or 'Ocorrencia registrada e aviso enviado pelo WhatsApp.')
        return

    logger.warning(
        'Ocorrencia registrada, mas envio de WhatsApp falhou.',
        extra={
            'aluno': aluno.nome,
            'aluno_id': aluno.id,
            'numero_usado': resultado_whatsapp.get('numero') or telefone,
            'tipo_ocorrencia': tipo_ocorrencia,
            'data_ocorrencia': str(data_ocorrencia),
            'erro_whatsapp': resultado_whatsapp.get('erro'),
            'resposta_whatsapp': resultado_whatsapp.get('resposta'),
            'status_code_whatsapp': resultado_whatsapp.get('status_code'),
        },
    )
    messages.warning(request, 'Ocorrencia registrada, mas o aviso nao foi enviado. Confira o telefone do aluno/responsavel.')


@login_required
@user_passes_test(grupo_digitadores, login_url='/')
def formulario_digitador(request):
    """View exclusiva para digitadores (apenas o formulario de ocorrencias)."""
    tipos_ocorrencia_permitidos = TIPOS_OCORRENCIA_PERMITIDOS

    # Recupera ultima data da sessao
    ultima_data_str = request.session.get('ultima_data_ocorrencia')
    if ultima_data_str:
        try:
            ultima_data = datetime.strptime(ultima_data_str, '%Y-%m-%d').date()
        except ValueError:
            ultima_data = timezone.now().date()
    else:
        ultima_data = timezone.now().date()

    # Recupera ultima turma da sessao
    ultima_turma_id = request.session.get('ultima_turma_ocorrencia_id')
    ultima_turma = None
    if ultima_turma_id:
        try:
            ultima_turma = Turma.objects.get(id=ultima_turma_id)
        except Turma.DoesNotExist:
            ultima_turma = None

    if request.method == 'POST':
        tipo_ocorrencia_postado = normalizar_tipo_ocorrencia(request.POST.get('tipo_ocorrencia'))
        if tipo_ocorrencia_postado not in tipos_ocorrencia_permitidos:
            messages.error(request, 'Tipo de ocorrencia invalido. Faltas devem ser registradas em Faltas de Alunos.')
            return redirect('formulario_digitador')

        form = RegistroOcorrenciaForm(request.POST)
        if form.is_valid():
            turma = form.cleaned_data['turma']
            numero = form.cleaned_data['numero_aluno']

            aluno = Aluno.objects.filter(turma=turma, numero=numero).first()
            if not aluno:
                messages.error(request, f'Aluno numero {numero} nao encontrado na turma {turma.nome}!')
                return render(request, 'ocorrencias/formulario_digitador.html', {
                    'form': form,
                    'ultimas_ocorrencias': RegistroOcorrenciaAluno.objects.select_related('aluno', 'aluno__turma').order_by('-data', '-horario_chegada')[:10]
                })

            ocorrencia = form.save(commit=False)

            # Pega o tipo de ocorrencia do formulario
            tipo_ocorrencia = tipo_ocorrencia_postado
            if tipo_ocorrencia not in tipos_ocorrencia_permitidos:
                messages.error(request, 'Tipo de ocorrencia invalido. Faltas devem ser registradas em Faltas de Alunos.')
                return redirect('formulario_digitador')
            enviar_reincidencia = request.POST.get('enviar_whatsapp_reincidencia') == 'sim'
            ocorrencia.tipo_ocorrencia = tipo_ocorrencia
            ocorrencia.faltou = False

            # Pega o turno do formulario
            ocorrencia.turno = form.cleaned_data.get('turno', 'manha')

            if ocorrencia.horario_chegada == '':
                ocorrencia.horario_chegada = None
            if ocorrencia.horario_contato == '':
                ocorrencia.horario_contato = None

            ocorrencia.aluno = aluno
            ocorrencia.registrado_por = request.user
            ja_existe_mesmo_tipo_antes = RegistroOcorrenciaAluno.objects.filter(
                aluno=aluno,
                data=ocorrencia.data,
                tipo_ocorrencia=tipo_ocorrencia,
            ).exists()
            try:
                with transaction.atomic():
                    ocorrencia.save()
            except IntegrityError:
                if tipo_ocorrencia in TIPOS_OCORRENCIA_REINCIDENCIA:
                    if enviar_reincidencia:
                        _enviar_aviso_ocorrencia_whatsapp_digitador(
                            request,
                            aluno,
                            tipo_ocorrencia,
                            ocorrencia.data,
                            mensagem_sucesso='Novo aviso de reincidencia enviado pelo WhatsApp.',
                        )
                    else:
                        messages.warning(request, 'Esta e a segunda ocorrencia deste tipo hoje. Deseja enviar novo aviso pelo WhatsApp?')
                messages.error(request, 'Ja existe um registro para este aluno nesta data com o mesmo tipo de ocorrencia. Nao e possivel duplicar sem alterar a restricao atual do banco.')
                return redirect('formulario_digitador')

            request.session['ultima_data_ocorrencia'] = ocorrencia.data.isoformat()
            request.session['ultima_turma_ocorrencia_id'] = turma.id
            request.session['ultima_turma_ocorrencia_nome'] = turma.nome

            messages.success(request, f'Ocorrencia registrada para {aluno.nome} (Turma {turma.nome}, No {numero})')
            if not ja_existe_mesmo_tipo_antes:
                _enviar_aviso_ocorrencia_whatsapp_digitador(request, aluno, tipo_ocorrencia, ocorrencia.data)
            return redirect('formulario_digitador')
        else:
            messages.error(request, 'Erro no formulario. Verifique os dados.')
    else:
        initial_data = {
            'data': ultima_data.isoformat(),
            'faltou': False,
        }
        form = RegistroOcorrenciaForm(initial=initial_data)
        if ultima_turma:
            form.fields['turma'].initial = ultima_turma

    ultimas_ocorrencias = RegistroOcorrenciaAluno.objects.select_related('aluno', 'aluno__turma').order_by('-data', '-horario_chegada')[:10]

    return render(request, 'ocorrencias/formulario_digitador.html', {
        'form': form,
        'ultimas_ocorrencias': ultimas_ocorrencias
    })

# =============================================================================
# IMPORTAÃ‡ÃƒO DE ALUNOS VIA EXCEL
# =============================================================================
from openpyxl import load_workbook
from ..models import Turma, Aluno

@login_required
@user_passes_test(pertence_ao_grupo_equipe_diretiva, login_url='/')
def importar_alunos_excel(request):
    if request.method == 'POST' and request.FILES.get('arquivo'):
        arquivo = request.FILES['arquivo']
        wb = load_workbook(arquivo)
        ws = wb.active

        contador = 0
        erros = 0

        for row in ws.iter_rows(min_row=2, values_only=True):  # Pula cabeÃ§alho
            try:
                turma_nome = str(row[0]).strip() if row[0] else ''  # Coluna A: Turma (ex: 3A)
                numero = int(row[1]) if row[1] else 0               # Coluna B: NÃºmero
                nome = str(row[2]).strip() if row[2] else ''        # Coluna C: Nome do aluno

                if turma_nome and numero and nome:
                    # Extrai a sÃ©rie do nome da turma (ex: 3A -> 3Âº Ano)
                    serie = f"{turma_nome[0]}Âº Ano"

                    # Busca ou cria a turma
                    turma, created = Turma.objects.get_or_create(
                        nome=turma_nome,
                        defaults={
                            'serie': serie,
                            'ano': 2026,
                            'ativa': True
                        }
                    )

                    # Cria o aluno
                    Aluno.objects.create(
                        nome=nome,
                        numero=numero,
                        turma=turma
                    )
                    contador += 1
                else:
                    erros += 1
            except Exception as e:
                erros += 1
                print(f"Erro na linha: {e}")

        messages.success(request, f'{contador} alunos importados com sucesso!')
        if erros > 0:
            messages.warning(request, f'{erros} linhas com erro foram ignoradas.')

        return redirect('controle_faltas_alunos')

    return render(request, 'importar_alunos.html')


@login_required
@user_passes_test(pertence_ao_grupo_equipe_diretiva, login_url='/')
def cadastrar_aluno(request):
    if request.method == 'POST':
        form = AlunoForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'âœ… Aluno cadastrado com sucesso!')
            return redirect('painel_equipe')
    else:
        form = AlunoForm()

    return render(request, 'cadastrar_aluno.html', {'form': form})


@login_required
@user_passes_test(pertence_ao_grupo_equipe_diretiva, login_url='/')
def listar_alunos(request):
    alunos = Aluno.objects.all().order_by('turma', 'numero')
    return render(request, 'listar_alunos.html', {'alunos': alunos})


@login_required
@user_passes_test(pertence_ao_grupo_equipe_diretiva, login_url='/')
def editar_aluno(request, aluno_id):
    aluno = get_object_or_404(Aluno, id=aluno_id)
    if request.method == 'POST':
        form = AlunoEditForm(request.POST, instance=aluno)
        if form.is_valid():
            form.save()
            messages.success(request, f'âœ… Aluno {aluno.nome} atualizado!')
            return redirect('listar_alunos')
    else:
        form = AlunoEditForm(instance=aluno)

    return render(request, 'editar_aluno.html', {'form': form, 'aluno': aluno})


# =============================================================================
# CADASTRO E EDIÃ‡ÃƒO DE ALUNOS (jÃ¡ existente)
# =============================================================================

@login_required
@user_passes_test(pertence_ao_grupo_equipe_diretiva, login_url='/')
def editar_aluno(request, aluno_id):
    aluno = get_object_or_404(Aluno, id=aluno_id)
    if request.method == 'POST':
        form = AlunoEditForm(request.POST, instance=aluno)
        if form.is_valid():
            form.save()
            messages.success(request, f'âœ… Aluno {aluno.nome} atualizado!')
            return redirect('listar_alunos')
    else:
        form = AlunoEditForm(instance=aluno)

    return render(request, 'editar_aluno.html', {'form': form, 'aluno': aluno})
