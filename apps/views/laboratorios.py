from .utilitarios import *

# =============================================================================
# GESTAO DE LABORATORIOS E EMPRESTIMOS
# =============================================================================

from ..models import Laboratorio, ItemEquipamento, AgendamentoLab, Emprestimo, Professor
from django.db.models import Max, Q
from django.urls import reverse
from datetime import datetime, timedelta


DIAS_SEMANA_LAB = ['Segunda-feira', 'Terca-feira', 'Quarta-feira', 'Quinta-feira', 'Sexta-feira']
HORARIOS_LAB = ['1', '2', '3', '4', '5', '6']
TURNOS_LAB = [
    ('manha', 'Manha'),
    ('tarde', 'Tarde'),
]
MAPA_DIAS_LAB = {
    'Monday': 'Segunda-feira',
    'Tuesday': 'Terca-feira',
    'Wednesday': 'Quarta-feira',
    'Thursday': 'Quinta-feira',
    'Friday': 'Sexta-feira',
    'Saturday': 'Sabado',
    'Sunday': 'Domingo',
}


def _laboratorios_agendaveis():
    return Laboratorio.objects.filter(ativo=True).order_by('nome')


def _laboratorios_filtrados(laboratorio_id):
    laboratorios = _laboratorios_agendaveis()
    if laboratorio_id:
        laboratorios = laboratorios.filter(id=laboratorio_id)
    return laboratorios


def _slots_cronograma(turno_filtro):
    turnos = [turno for turno in TURNOS_LAB if turno[0] == turno_filtro] if turno_filtro in ['manha', 'tarde'] else TURNOS_LAB
    return [
        {
            'turno': turno,
            'turno_label': turno_label,
            'horario': horario,
            'horario_label': f'Aula{horario}',
        }
        for turno, turno_label in turnos
        for horario in HORARIOS_LAB
    ]


def _inicio_semana(data):
    return data - timedelta(days=data.weekday())


def _data_inicio_cronograma(request, laboratorio_id='', turno_filtro=''):
    data_inicio = request.GET.get('data_inicio')
    if data_inicio:
        return _inicio_semana(datetime.strptime(data_inicio, '%Y-%m-%d').date())

    agendamentos = AgendamentoLab.objects.all()
    if laboratorio_id:
        agendamentos = agendamentos.filter(laboratorio_id=laboratorio_id)
    if turno_filtro in ['manha', 'tarde']:
        agendamentos = agendamentos.filter(turno=turno_filtro)

    data_recente = agendamentos.aggregate(data=Max('data'))['data']
    if data_recente:
        return _inicio_semana(data_recente)

    return _inicio_semana(datetime.now().date())


def _agendamentos_por_chave(agendamentos):
    agendamentos_dict = {}
    for agendamento in agendamentos:
        dia_en = agendamento.data.strftime('%A')
        dia_pt = MAPA_DIAS_LAB.get(dia_en, dia_en)
        agendamentos_dict[(agendamento.laboratorio.id, dia_pt, agendamento.turno, agendamento.horario)] = agendamento
    return agendamentos_dict


def _agendamentos_cards(agendamentos):
    cards = []
    turno_labels = dict(TURNOS_LAB)
    for agendamento in agendamentos:
        dia_en = agendamento.data.strftime('%A')
        cards.append({
            'agendamento': agendamento,
            'dia': MAPA_DIAS_LAB.get(dia_en, dia_en),
            'turno_codigo': agendamento.turno,
            'turno_label': turno_labels.get(agendamento.turno, agendamento.turno),
            'horario_label': f'Aula{agendamento.horario}',
        })
    return cards


def _grupos_turno_cronograma(turno_filtro, agendamentos_cards):
    turnos = [turno for turno in TURNOS_LAB if turno[0] == turno_filtro] if turno_filtro in ['manha', 'tarde'] else TURNOS_LAB
    estilos = {
        'manha': ('bg-primary', 'text-white'),
        'tarde': ('bg-warning', 'text-dark'),
    }
    grupos = []
    for turno, turno_label in turnos:
        bg_class, text_class = estilos[turno]
        grupos.append({
            'turno': turno,
            'turno_label': turno_label,
            'titulo': turno_label.upper(),
            'bg_class': bg_class,
            'text_class': text_class,
            'slots': _slots_cronograma(turno),
            'cards': [card for card in agendamentos_cards if card['turno_codigo'] == turno],
        })
    return grupos


def _horarios_laboratorio_professor(user, turno_filtro=''):
    if turno_filtro not in ['manha', 'tarde']:
        turno_filtro = ''

    professor = Professor.objects.filter(usuario=user, ativo=True).first()
    if not professor:
        return {
            'professor': None,
            'agendamentos': [],
            'semana_inicio': None,
            'semana_fim': None,
            'turno_filtro': turno_filtro,
        }

    hoje = timezone.localdate()
    semana_inicio = _inicio_semana(hoje)
    semana_fim = semana_inicio + timedelta(days=6)

    base = AgendamentoLab.objects.filter(professor=professor)
    agendamentos = base.filter(data__range=(semana_inicio, semana_fim))

    if not agendamentos.exists():
        data_recente = base.aggregate(data=Max('data'))['data']
        if data_recente:
            semana_inicio = _inicio_semana(data_recente)
            semana_fim = semana_inicio + timedelta(days=6)
            agendamentos = base.filter(data__range=(semana_inicio, semana_fim))

    if turno_filtro:
        agendamentos = agendamentos.filter(turno=turno_filtro)

    agendamentos = agendamentos.select_related('laboratorio', 'turma').order_by(
        'data', 'turno', 'horario', 'laboratorio__nome'
    )
    turno_labels = dict(TURNOS_LAB)
    itens = []
    for agendamento in agendamentos:
        dia_en = agendamento.data.strftime('%A')
        itens.append({
            'agendamento': agendamento,
            'dia': MAPA_DIAS_LAB.get(dia_en, dia_en),
            'turno_codigo': agendamento.turno,
            'turno': turno_labels.get(agendamento.turno, agendamento.turno),
            'aula': f'Aula{agendamento.horario}',
        })

    return {
        'professor': professor,
        'agendamentos': itens,
        'semana_inicio': semana_inicio,
        'semana_fim': semana_fim,
        'turno_filtro': turno_filtro,
    }


@login_required
def meus_horarios_laboratorio(request):
    context = {
        'horarios_laboratorio_professor': _horarios_laboratorio_professor(
            request.user,
            request.GET.get('turno', ''),
        ),
    }
    return render(request, 'laboratorios/meus_horarios.html', context)

@login_required
@user_passes_test(pertence_ao_grupo_equipe_diretiva, login_url='/')
def listar_laboratorios(request):
    """Lista todos os laboratorios"""
    laboratorios = Laboratorio.objects.filter(ativo=True)

    # Contagem de disponiveis para itinerantes
    for lab in laboratorios:
        if lab.tipo == 'itinerante':
            lab.disponiveis = ItemEquipamento.objects.filter(
                laboratorio=lab,
                disponivel=True
            ).count()

    return render(request, 'laboratorios/listar.html', {'laboratorios': laboratorios})

@login_required
@user_passes_test(pertence_ao_grupo_equipe_diretiva, login_url='/')
def excluir_agendamento(request, agendamento_id):
    from django.shortcuts import get_object_or_404, redirect
    from django.contrib import messages
    from ..models import AgendamentoLab   # ajuste o nome do modelo se necessario

    agendamento = get_object_or_404(AgendamentoLab, id=agendamento_id)
    agendamento.delete()
    messages.success(request, "Agendamento removido com sucesso!")
    data_inicio = request.GET.get('data_inicio', '')
    turno = request.GET.get('turno', '')
    laboratorio = request.GET.get('laboratorio', '')
    url = reverse('cronograma_semanal')
    params = []
    if data_inicio:
        params.append(f'data_inicio={data_inicio}')
    if turno:
        params.append(f'turno={turno}')
    if laboratorio:
        params.append(f'laboratorio={laboratorio}')
    if params:
        url = f'{url}?{"&".join(params)}'
    return redirect(url)

@login_required
@user_passes_test(pertence_ao_grupo_equipe_diretiva, login_url='/')
def agendamento_lab(request, lab_id):
    """Agendar laboratorio por aula (Labs 01 a 05)."""
    laboratorio = get_object_or_404(Laboratorio, id=lab_id, ativo=True)

    if request.method == 'POST':
        data = request.POST.get('data')
        horario = request.POST.get('horario')
        turno = request.POST.get('turno')
        professor_id = request.POST.get('professor')
        disciplina = request.POST.get('disciplina')
        turma_id = request.POST.get('turma')
        observacao = request.POST.get('observacao', '')

        # 1. Verificar se o mesmo professor ja tem agendamento neste horario (em qualquer lab)
        conflito_professor = AgendamentoLab.objects.filter(
            professor_id=professor_id,
            data=data,
            horario=horario,
            turno=turno
        ).exists()

        if conflito_professor:
            messages.error(request, 'Este professor ja possui agendamento neste horario em outro laboratorio!')
            # Buscar dados para renderizar o formulario novamente
            from ..models import Professor, Turma
            professores = Professor.objects.filter(ativo=True)
            turmas = Turma.objects.filter(ativa=True)
            return render(request, 'laboratorios/agendar.html', {
                'laboratorio': laboratorio,
                'laboratorios': _laboratorios_agendaveis(),
                'professores': professores,
                'turmas': turmas,
            })

        # 2. Verificar se o laboratorio ja esta reservado neste horario
        conflito_lab = AgendamentoLab.objects.filter(
            laboratorio=laboratorio,
            data=data,
            horario=horario,
            turno=turno
        ).exists()

        if conflito_lab:
            messages.error(request, f'{laboratorio.nome} ja esta reservado neste horario!')
            from ..models import Professor, Turma
            professores = Professor.objects.filter(ativo=True)
            turmas = Turma.objects.filter(ativa=True)
            return render(request, 'laboratorios/agendar.html', {
                'laboratorio': laboratorio,
                'laboratorios': _laboratorios_agendaveis(),
                'professores': professores,
                'turmas': turmas,
            })

        # 3. Salvar agendamento
        from ..models import Professor, Turma
        AgendamentoLab.objects.create(
            laboratorio=laboratorio,
            data=data,
            horario=horario,
            turno=turno,
            professor_id=professor_id,
            disciplina=disciplina,
            turma_id=turma_id,
            observacao=observacao,
            registrado_por=request.user
        )
        messages.success(request, f'{laboratorio.nome} agendado com sucesso!')
        return redirect('listar_laboratorios')

    # GET - mostrar formulario
    from ..models import Professor, Turma
    professores = Professor.objects.filter(ativo=True)
    turmas = Turma.objects.filter(ativa=True)

    return render(request, 'laboratorios/agendar.html', {
        'laboratorio': laboratorio,
        'laboratorios': _laboratorios_agendaveis(),
        'professores': professores,
        'turmas': turmas,
    })


@login_required
@login_required(login_url='/')
def cronograma_semanal(request):
    """Exibe o cronograma semanal dos laboratorios, com filtro opcional por turno."""
    # Capturar o filtro de turno (vindo do template)
    turno_filtro = request.GET.get('turno', '')
    laboratorio_filtro = request.GET.get('laboratorio', '')
    data_inicio = _data_inicio_cronograma(request, laboratorio_filtro, turno_filtro)
    data_fim = data_inicio + timedelta(days=6)

    laboratorios_todos = _laboratorios_agendaveis()
    laboratorios = _laboratorios_filtrados(laboratorio_filtro)

    agendamentos = AgendamentoLab.objects.filter(
        data__gte=data_inicio,
        data__lte=data_fim,
        laboratorio__in=laboratorios,
    )

    # Aplicar filtro de turno se fornecido
    if turno_filtro in ['manha', 'tarde']:
        agendamentos = agendamentos.filter(turno=turno_filtro)

    agendamentos = agendamentos.select_related('laboratorio', 'professor', 'turma').order_by(
        'laboratorio__nome', 'data', 'horario'
    )
    agendamentos_lista = list(agendamentos)
    agendamentos_cards = _agendamentos_cards(agendamentos_lista)

    context = {
        'dias_semana': DIAS_SEMANA_LAB,
        'horarios': HORARIOS_LAB,
        'slots_cronograma': _slots_cronograma(turno_filtro),
        'agendamentos_dict': _agendamentos_por_chave(agendamentos_lista),
        'agendamentos_cards': agendamentos_cards,
        'grupos_turno_cronograma': _grupos_turno_cronograma(turno_filtro, agendamentos_cards),
        'laboratorios': laboratorios,
        'laboratorios_todos': laboratorios_todos,
        'data_inicio': data_inicio,
        'data_fim': data_fim,
        'turno_filtro': turno_filtro,
        'laboratorio_filtro': laboratorio_filtro,
    }
    return render(request, 'laboratorios/cronograma.html', context)

@login_required
@user_passes_test(pertence_ao_grupo_equipe_diretiva, login_url='/')
def cronograma_print(request):
    """Exibe o cronograma semanal para impressao."""
    # 🔥 CAPTURAR TURNO (FALTAVA ISSO)
    turno_filtro = request.GET.get('turno', '')
    laboratorio_filtro = request.GET.get('laboratorio', '')
    data_inicio = _data_inicio_cronograma(request, laboratorio_filtro, turno_filtro)
    data_fim = data_inicio + timedelta(days=6)

    laboratorios_todos = _laboratorios_agendaveis()
    laboratorios = _laboratorios_filtrados(laboratorio_filtro)

    # 📊 Buscar agendamentos
    agendamentos = AgendamentoLab.objects.filter(
        data__gte=data_inicio,
        data__lte=data_fim,
        laboratorio__in=laboratorios,
    )

    # 🔥 APLICAR FILTRO (FALTAVA ISSO)
    if turno_filtro in ['manha', 'tarde']:
        agendamentos = agendamentos.filter(turno=turno_filtro)

    agendamentos = agendamentos.select_related('laboratorio', 'professor', 'turma').order_by(
        'laboratorio__nome', 'data', 'horario'
    )
    agendamentos_lista = list(agendamentos)
    agendamentos_cards = _agendamentos_cards(agendamentos_lista)

    # 📤 Contexto
    context = {
        'dias_semana': DIAS_SEMANA_LAB,
        'horarios': HORARIOS_LAB,
        'slots_cronograma': _slots_cronograma(turno_filtro),
        'agendamentos_dict': _agendamentos_por_chave(agendamentos_lista),
        'agendamentos_cards': agendamentos_cards,
        'grupos_turno_cronograma': _grupos_turno_cronograma(turno_filtro, agendamentos_cards),
        'laboratorios': laboratorios,
        'laboratorios_todos': laboratorios_todos,
        'data_inicio': data_inicio,
        'data_fim': data_fim,
        'turno_filtro': turno_filtro,
        'laboratorio_filtro': laboratorio_filtro,
    }

    return render(request, 'laboratorios/cronograma_print.html', context)

    import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from django.http import HttpResponse

@login_required
@user_passes_test(pertence_ao_grupo_equipe_diretiva, login_url='/')
def exportar_cronograma_excel(request):
    from ..models import Laboratorio, AgendamentoLab
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    # Captura o filtro de turno (vem do template)
    turno_filtro = request.GET.get('turno', '')
    laboratorio_filtro = request.GET.get('laboratorio', '')
    data_inicio = _data_inicio_cronograma(request, laboratorio_filtro, turno_filtro)
    data_fim = data_inicio + timedelta(days=6)

    laboratorios = _laboratorios_filtrados(laboratorio_filtro)
    dias_semana = DIAS_SEMANA_LAB
    horarios = HORARIOS_LAB

    # Buscar agendamentos da semana (com filtro de turno)
    agendamentos = AgendamentoLab.objects.filter(
        data__gte=data_inicio,
        data__lte=data_fim,
        laboratorio__in=laboratorios,
    )
    if turno_filtro in ['manha', 'tarde']:
        agendamentos = agendamentos.filter(turno=turno_filtro)
    agendamentos = agendamentos.select_related('laboratorio', 'professor', 'turma')

    # Dicionario para acesso rapido
    agendamentos_dict = _agendamentos_por_chave(agendamentos)

    wb = Workbook()
    wb.remove(wb.active)

    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    for lab in laboratorios:
        ws = wb.create_sheet(title=lab.nome)

        # Cabecalho: dias da semana (uma coluna por dia)
        for idx, dia in enumerate(dias_semana, start=2):
            cell = ws.cell(row=1, column=idx, value=dia)
            cell.font = Font(bold=True, size=11)
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
            cell.border = thin_border

        # Agora, para cada horario, mesclar 3 linhas na coluna A e preencher as outras colunas
        row_offset = 2  # comeca na linha 2 (apos cabecalho)
        for horario_idx, horario in enumerate(horarios):
            base_row = row_offset + horario_idx * 3

            # Mescla as 3 linhas na coluna A (horario)
            ws.merge_cells(start_row=base_row, start_column=1, end_row=base_row+2, end_column=1)
            cell_horario = ws.cell(row=base_row, column=1, value=f"Aula{horario}")
            cell_horario.font = Font(bold=True, size=11)
            cell_horario.alignment = Alignment(horizontal='center', vertical='center')
            cell_horario.border = thin_border

            # Para cada dia da semana
            for col_idx, dia in enumerate(dias_semana, start=2):
                key = (lab.id, dia, 'manha' if turno_filtro == 'manha' else 'tarde' if turno_filtro == 'tarde' else 'manha', horario)
                a = agendamentos_dict.get(key)
                if not a and turno_filtro not in ['manha', 'tarde']:
                    a = agendamentos_dict.get((lab.id, dia, 'tarde', horario))

                # Linha da disciplina
                ws.cell(row=base_row, column=col_idx, value=a.disciplina if a else "").border = thin_border
                ws.cell(row=base_row, column=col_idx).alignment = Alignment(horizontal='center', vertical='center')
                # Linha da turma
                ws.cell(row=base_row+1, column=col_idx, value=a.turma.nome if a else "").border = thin_border
                ws.cell(row=base_row+1, column=col_idx).alignment = Alignment(horizontal='center', vertical='center')
                # Linha do professor
                ws.cell(row=base_row + 2, column=col_idx,
                        value=a.professor.nome_abreviado if a else "").border = thin_border
                ws.cell(row=base_row+2, column=col_idx).alignment = Alignment(horizontal='center', vertical='center')

        # Ajustar larguras
        ws.column_dimensions['A'].width = 12
        for col in range(2, 2 + len(dias_semana)):
            col_letter = get_column_letter(col)
            ws.column_dimensions[col_letter].width = 28

        # Ajustar altura das linhas (cada bloco de 3 linhas = 60, cada linha 20)
        for row in range(row_offset, row_offset + len(horarios)*3):
            ws.row_dimensions[row].height = 20

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="cronograma_{data_inicio.strftime("%Y%m%d")}.xlsx"'
    wb.save(response)
    return response


@login_required
@user_passes_test(pertence_ao_grupo_equipe_diretiva, login_url='/')
def emprestimo_equipamento(request, lab_id):
    """Registrar emprestimo de equipamentos (Lab 04 ou Lab 05)"""
    laboratorio = get_object_or_404(Laboratorio, id=lab_id, tipo='itinerante')

    # Equipamentos disponiveis
    equipamentos_disponiveis = ItemEquipamento.objects.filter(
        laboratorio=laboratorio,
        disponivel=True
    )

    if request.method == 'POST':
        tipo_emprestimo = request.POST.get('tipo_emprestimo')
        quantidade = int(request.POST.get('quantidade', 0))
        data_prevista = request.POST.get('data_prevista_devolucao')
        motivo = request.POST.get('motivo')
        aulas_utilizacao = request.POST.get('aulas_utilizacao', '')
        observacao = request.POST.get('observacao', '')

        # Pegar equipamentos selecionados (checkbox)
        itens_selecionados = request.POST.getlist('itens')

        # Validar
        if len(itens_selecionados) != quantidade:
            messages.error(request, 'Selecione a quantidade correta de equipamentos!')
        else:
            # Criar emprestimo
            emprestimo = Emprestimo.objects.create(
                tipo_emprestimo=tipo_emprestimo,
                laboratorio=laboratorio,
                quantidade=quantidade,
                data_prevista_devolucao=data_prevista,
                motivo=motivo,
                aulas_utilizacao=aulas_utilizacao,
                observacao=observacao,
                registrado_por=request.user
            )

            # Adicionar itens e marcar como indisponiveis
            for item_id in itens_selecionados:
                item = ItemEquipamento.objects.get(id=item_id)
                emprestimo.itens.add(item)
                item.disponivel = False
                item.save()

            # Se for professor
            if tipo_emprestimo == 'professor':
                emprestimo.professor_id = request.POST.get('professor_id')
            elif tipo_emprestimo == 'aluno':
                emprestimo.aluno_id = request.POST.get('aluno_id')
            elif tipo_emprestimo == 'turma':
                emprestimo.turma_id = request.POST.get('turma_id')

            emprestimo.save()
            messages.success(request, f'Emprestimo registrado! {quantidade} equipamento(s) emprestado(s).')
            return redirect('listar_laboratorios')

    # GET - mostrar formulario
    from ..models import Professor, Aluno, Turma
    professores = Professor.objects.filter(ativo=True)
    alunos = Aluno.objects.filter(ativo=True)
    turmas = Turma.objects.filter(ativa=True)

    return render(request, 'laboratorios/emprestimo.html', {
        'laboratorio': laboratorio,
        'equipamentos': equipamentos_disponiveis,
        'professores': professores,
        'alunos': alunos,
        'turmas': turmas,
    })


@login_required
@user_passes_test(pertence_ao_grupo_equipe_diretiva, login_url='/')
def devolver_equipamento(request, emprestimo_id):
    """Registrar devolucao de equipamentos"""
    emprestimo = get_object_or_404(Emprestimo, id=emprestimo_id)

    if request.method == 'POST':
        from django.utils import timezone
        emprestimo.data_devolucao = timezone.now()
        emprestimo.status = 'devolvido'
        emprestimo.save()

        # Marcar equipamentos como disponiveis novamente
        for item in emprestimo.itens.all():
            item.disponivel = True
            item.save()

        messages.success(request, 'Equipamentos devolvidos com sucesso!')
        return redirect('listar_emprestimos')

    return render(request, 'laboratorios/devolver.html', {'emprestimo': emprestimo})


@login_required
@user_passes_test(pertence_ao_grupo_equipe_diretiva, login_url='/')
def listar_emprestimos(request):
    """Lista todos os emprestimos ativos"""
    emprestimos = Emprestimo.objects.filter(status='emprestado').order_by('data_prevista_devolucao')
    return render(request, 'laboratorios/lista_emprestimos.html', {'emprestimos': emprestimos})
